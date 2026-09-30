"""
Unit tests for Corporate Excel S&OP Workbook generator.
Verifies live formula syntax and workbook structure with openpyxl.
"""
import os
import openpyxl
import pytest
from src.reporting.excel_snop_builder import CorporateSnOpExcelBuilder


def test_excel_snop_builder_structure_and_formulas(tmp_path):
    """Valida la presencia de fórmulas encadenadas vivas en la planilla generada."""
    out_file = str(tmp_path / "test_snop_model.xlsx")
    CorporateSnOpExcelBuilder.build_snop_workbook(filepath=out_file)

    assert os.path.exists(out_file)
    wb = openpyxl.load_workbook(out_file, data_only=False)

    # 1. Verifica pestañas
    sheet_names = wb.sheetnames
    assert "Plan_Compras_SOP" in sheet_names
    assert "Auditoria_Forecast" in sheet_names
    assert "Monitor_Caducidad_FEFO" in sheet_names

    ws1 = wb["Plan_Compras_SOP"]

    # 2. Verifica fila de encabezados
    headers = [cell.value for cell in ws1[1]]
    assert "SKU" in headers
    assert "Requerimiento Neto" in headers
    assert "Orden Sugerida" in headers
    assert "Estatus S&OP" in headers

    # 3. Verifica fórmulas vivas en la fila 2
    row2_j = str(ws1["J2"].value)  # Posición Neta
    assert row2_j.startswith("=")
    assert "F2" in row2_j and "G2" in row2_j and "H2" in row2_j

    row2_k = str(ws1["K2"].value)  # Requerimiento Neto
    assert row2_k.startswith("=")
    assert "MAX" in row2_k

    row2_m = str(ws1["M2"].value)  # Orden Sugerida
    assert row2_m.startswith("=")
    assert "IF" in row2_m and "MAX" in row2_m

    row2_o = str(ws1["O2"].value)  # Inversión Sugerida
    assert row2_o.startswith("=")
    assert "M2*N2" in row2_o

    row2_p = str(ws1["P2"].value)  # Cobertura Semanas
    assert row2_p.startswith("=")
    assert "ROUND" in row2_p

    row2_q = str(ws1["Q2"].value)  # Estatus S&OP
    assert row2_q.startswith("=")
    assert "APROBAR COMPRA" in row2_q

    # 4. Verifica fórmulas de totales
    max_r = ws1.max_row
    total_val = str(ws1.cell(row=max_r, column=15).value)
    assert total_val.startswith("=SUM(")
